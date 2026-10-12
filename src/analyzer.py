"""Access risk analyzer.

1. Hashes and validates the inputs (stops on any validation error)
2. Joins application accounts to the HR roster
3. Flags terminated, dormant, orphaned, excess-privilege and SoD findings
4. Assigns severity and maps each finding to its control
5. Writes output/findings_report.xlsx (sanitized) and output/findings.csv
6. Compares results to data/planted_issues.csv and reports detection accuracy

Usage:  python src/analyzer.py [--hr PATH] [--accounts PATH] [--sod PATH]
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config as C
from common import account_ref, audit, sanitize_value, sha256_file, write_sanitized_csv
from validation import ValidationError, load_and_validate

KEY_COLS = ["issue_type", "username", "application", "detail"]
FINDING_COLUMNS = ["issue_type", "severity", "control_id", "control_requirement", "account_ref",
                   "username", "employee_id", "job_title", "department", "store_location",
                   "manager_id", "employee_status", "application", "detail", "last_login_date",
                   "days_since_login", "evidence", "recommended_action"]


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------
def hash_inputs(paths: list[Path]) -> pd.DataFrame:
    """SHA-256 every input before it is read, and log each hash."""
    rows = []
    for p in paths:
        if not p.exists():
            raise ValidationError(f"Input validation failed:\n  - {p.name}: file not found at {p}")
        digest = sha256_file(p)
        rows.append({"file": p.name, "sha256": digest, "bytes": p.stat().st_size})
        audit(f"analyzer: hashed input {p.name}", "-", digest)
    return pd.DataFrame(rows)


def run_analysis(hr_path: Path = C.HR_FILE, accounts_path: Path = C.ACCOUNTS_FILE,
                 sod_path: Path = C.SOD_FILE):
    """Hash -> validate -> analyze. Shared by the packets script and dashboard."""
    hashes = hash_inputs([hr_path, accounts_path, sod_path])
    hr, accts, sod = load_and_validate(hr_path, accounts_path, sod_path)
    audit("analyzer: validation passed", len(hr) + len(accts) + len(sod))
    findings, accounts, warnings = analyze(hr, accts, sod)
    return hr, accts, sod, accounts, findings, warnings, hashes


# ---------------------------------------------------------------------------
# Analysis
# ---------------------------------------------------------------------------
def build_accounts(hr: pd.DataFrame, accts: pd.DataFrame) -> pd.DataFrame:
    """Collapse entitlement rows to one row per account and join HR."""
    acc = (accts.groupby(["application", "username"], as_index=False)
           .agg(employee_id=("employee_id", "first"),
                last_login_date=("last_login_date", "max"),
                entitlements=("entitlement", lambda s: sorted(set(s)))))
    hr_cols = ["employee_id", "name", "job_title", "department", "manager_id",
               "store_location", "status", "termination_date"]
    acc = acc.merge(hr[hr_cols], on="employee_id", how="left", indicator="hr_match")
    acc["has_hr_record"] = (acc["hr_match"] == "both") & (acc["employee_id"] != "")
    acc = acc.drop(columns="hr_match")
    for col in hr_cols[1:]:
        acc[col] = acc[col].fillna("")
    as_of = pd.Timestamp(C.AS_OF_DATE)
    acc["days_since_login"] = (as_of - pd.to_datetime(acc["last_login_date"])).dt.days
    acc["account_ref"] = [account_ref(u, a) for u, a in zip(acc["username"], acc["application"])]
    return acc


def _finding(acct: pd.Series, issue_type: str, severity: str, detail: str, evidence: str,
             action: str) -> dict:
    control_id, control_desc = C.CONTROLS[issue_type]
    return {
        "issue_type": issue_type, "severity": severity, "control_id": control_id,
        "control_requirement": control_desc,
        "account_ref": acct["account_ref"], "username": acct["username"],
        "employee_id": acct["employee_id"], "job_title": acct["job_title"],
        "department": acct["department"], "store_location": acct["store_location"],
        "manager_id": acct["manager_id"], "employee_status": acct["status"] or "no HR record",
        "application": acct["application"], "detail": detail,
        "last_login_date": acct["last_login_date"], "days_since_login": int(acct["days_since_login"]),
        "evidence": evidence, "recommended_action": action,
    }


def analyze(hr: pd.DataFrame, accts: pd.DataFrame, sod: pd.DataFrame):
    acc = build_accounts(hr, accts)
    out: list[dict] = []

    for _, a in acc.iterrows():
        ents = set(a["entitlements"])
        financial = a["application"] in C.FINANCIAL_APPS
        privileged = bool(ents & C.MANAGER_LEVEL_ENTITLEMENTS)

        # Orphaned: no matching HR record (NIST AC-2)
        if not a["has_hr_record"]:
            why = "blank employee_id" if a["employee_id"] == "" else "employee_id not in HR roster"
            sev = "High" if (financial or privileged) else "Medium"
            out.append(_finding(a, C.ORPHAN, sev, "",
                                f"No HR owner ({why}); {len(ents)} entitlement(s)",
                                "Disable now; re-enable only after an owner is assigned in HR"))
            # Orphans can still be dormant, so fall through to the dormancy check.

        # Terminated with active access (PCI-DSS 8.2.5)
        if a["status"] == "terminated":
            term = pd.Timestamp(a["termination_date"])
            evidence = f"Terminated {a['termination_date']}; account still enabled"
            if pd.Timestamp(a["last_login_date"]) > term:
                evidence += f"; LOGIN AFTER TERMINATION on {a['last_login_date']}"
            out.append(_finding(a, C.TERMINATED, "High", "", evidence,
                                "Disable immediately and review activity since termination date"))

        # Dormant > 90 days (PCI-DSS 8.2.6)
        if a["days_since_login"] > C.DORMANCY_DAYS:
            sev = "Medium" if (financial or privileged or a["days_since_login"] > 180) else "Low"
            out.append(_finding(a, C.DORMANT, sev, "",
                                f"No login for {a['days_since_login']} days (limit {C.DORMANCY_DAYS})",
                                "Disable; manager confirms business need before re-enabling"))

        # Excess privilege: frontline role holding manager-level rights (NIST AC-6)
        if a["job_title"] in C.FRONTLINE_TITLES:
            for ent in sorted(ents & C.MANAGER_LEVEL_ENTITLEMENTS):
                sev = "High" if ent in {"POS: Store Manager Admin",
                                        "POS: Approve Cash Reconciliation"} else "Medium"
                out.append(_finding(a, C.EXCESS, sev, ent,
                                    f"'{a['job_title']}' holds manager-level entitlement",
                                    f"Remove '{ent}'; grant time-bound access if coverage is needed"))

    # SoD: evaluate per identity across ALL applications, so a rule that spans
    # two apps is still caught. Orphans are keyed by username.
    acc["identity"] = acc["employee_id"].where(acc["has_hr_record"], "ORPHAN:" + acc["username"])
    for identity, grp in acc.groupby("identity"):
        ent_to_acct = {e: row for _, row in grp.iterrows() for e in row["entitlements"]}
        for _, rule in sod.iterrows():
            a_ent, b_ent = rule["entitlement_a"], rule["entitlement_b"]
            if a_ent in ent_to_acct and b_ent in ent_to_acct:
                a = ent_to_acct[a_ent].copy()
                b_app = ent_to_acct[b_ent]["application"]
                if b_app != a["application"]:
                    a["application"] = f"{a['application']} + {b_app}"
                out.append(_finding(a, C.SOD, rule["severity"], rule["rule_id"],
                                    f"Holds both '{a_ent}' and '{b_ent}'. {rule['risk_description']}",
                                    f"Remove one of the two; if impossible, document a "
                                    f"compensating control (e.g. secondary review)"))

    findings = pd.DataFrame(out)
    if findings.empty:
        findings = pd.DataFrame(columns=FINDING_COLUMNS)
    findings["sev_rank"] = findings["severity"].map(C.SEVERITY_ORDER)
    findings = (findings.sort_values(["sev_rank", "issue_type", "application", "employee_id"])
                .drop(columns="sev_rank").reset_index(drop=True))
    findings.insert(0, "finding_id", [f"F-{i:04d}" for i in range(1, len(findings) + 1)])

    warnings = data_quality_warnings(hr)
    return findings, acc, warnings


def data_quality_warnings(hr: pd.DataFrame) -> list[dict]:
    """HR issues that don't violate a control but affect the review."""
    warnings = []
    status = hr.set_index("employee_id")["status"]
    active = hr[hr["status"] == "active"]
    to_term = active[active["manager_id"].map(status) == "terminated"]
    for mgr_id, grp in to_term.groupby("manager_id"):
        warnings.append({
            "warning": "Active employees report to a terminated manager",
            "manager_id": mgr_id, "affected_employees": len(grp),
            "impact": "Packet is escalated to the terminated manager's own manager",
            "fix": "HR should update manager_id for these employees",
        })
    return warnings


# ---------------------------------------------------------------------------
# Detection accuracy
# ---------------------------------------------------------------------------
def detection_accuracy(findings: pd.DataFrame, key_path: Path = C.ANSWER_KEY_FILE):
    if not key_path.exists():
        return None, None
    key = pd.read_csv(key_path, dtype=str, keep_default_na=False)
    found = findings[KEY_COLS].drop_duplicates()
    merged = found.merge(key[KEY_COLS].drop_duplicates(), how="outer", indicator=True)
    merged["result"] = merged["_merge"].map(
        {"both": "True positive", "left_only": "False positive", "right_only": "Missed (false negative)"})
    rows = []
    for t in C.FINDING_TYPES:
        sub = merged[merged["issue_type"] == t]
        rows.append({"issue_type": t,
                     "planted": int((sub["_merge"] != "left_only").sum()),
                     "detected": int((sub["_merge"] != "right_only").sum()),
                     "true_positives": int((sub["_merge"] == "both").sum()),
                     "false_positives": int((sub["_merge"] == "left_only").sum()),
                     "missed": int((sub["_merge"] == "right_only").sum())})
    summary = pd.DataFrame(rows)
    detail = merged.drop(columns="_merge")
    return summary, detail


# ---------------------------------------------------------------------------
# Excel report
# ---------------------------------------------------------------------------
FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
SEV_FILL = {"High": PatternFill("solid", start_color="F8CBAD"),
            "Medium": PatternFill("solid", start_color="FFE699"),
            "Low": PatternFill("solid", start_color="E2EFDA")}
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))


def _banner(ws, title: str, width: int) -> int:
    """Confidentiality marking in-sheet (row 1) and in the printed header."""
    ws["A1"] = C.CONFIDENTIAL.upper()
    ws["A1"].font = Font(name=FONT, bold=True, color="C00000", size=11)
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, bold=True, size=14)
    ws.oddHeader.center.text = C.CONFIDENTIAL
    ws.oddFooter.center.text = "Page &P of &N"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return 4


def _table(ws, df: pd.DataFrame, start_row: int, sev_col: str | None = None,
           widths: dict | None = None) -> int:
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start_row, column=j, value=sanitize_value(str(col)))
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, val in enumerate(row, 1):
            if isinstance(val, (list, tuple, set)):
                val = ", ".join(map(str, val))
            c = ws.cell(row=i, column=j, value=sanitize_value(val))
            c.font = Font(name=FONT, size=10)
            c.border = THIN
            c.alignment = Alignment(vertical="top", wrap_text=False)
        if sev_col:
            sev = getattr(row, sev_col)
            if sev in SEV_FILL:
                for j in range(1, len(df.columns) + 1):
                    ws.cell(row=i, column=j).fill = SEV_FILL[sev]
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col)
        if w is None:
            longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
            w = min(max(10, longest + 2), 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    last = start_row + len(df)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{last}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return last


def write_report(findings, accounts, warnings, hashes, acc_summary, acc_detail, path: Path) -> str:
    wb = Workbook()
    run_ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    n_find = len(findings)
    f_last = 5 + n_find  # Findings header row is 5

    # --- Summary (formulas reference the Findings sheet so it recalculates)
    ws = wb.active
    ws.title = "Summary"
    r = _banner(ws, "Quarterly Access Review: Findings Report", 6)
    meta = [("Company", C.COMPANY), ("Review as-of date", C.AS_OF_DATE.isoformat()),
            ("Report generated", run_ts),
            ("Accounts reviewed", len(accounts)),
            ("Identities reviewed (HR-matched)", int(accounts.loc[accounts["has_hr_record"], "employee_id"].nunique())),
            ("Entitlement rows reviewed", int(accounts["entitlements"].map(len).sum()))]
    for label, val in meta:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value=sanitize_value(val)).font = Font(name=FONT)
        r += 1
    r += 1
    head = ["Finding type", "Control", "High", "Medium", "Low", "Total"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
    first_data = r + 1
    rng_type = f"Findings!$B$6:$B${max(f_last, 6)}"
    rng_sev = f"Findings!$C$6:$C${max(f_last, 6)}"
    for t in C.FINDING_TYPES:
        r += 1
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=C.CONTROLS[t][0])
        for j, sev in zip((3, 4, 5), ("High", "Medium", "Low")):
            ws.cell(row=r, column=j, value=f'=COUNTIFS({rng_type},$A{r},{rng_sev},"{sev}")')
        ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})")
        for j in range(1, 7):
            ws.cell(row=r, column=j).font = Font(name=FONT)
    r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, bold=True)
    for j in range(3, 7):
        col = get_column_letter(j)
        c = ws.cell(row=r, column=j, value=f"=SUM({col}{first_data}:{col}{r - 1})")
        c.font = Font(name=FONT, bold=True)
    r += 2
    ws.cell(row=r, column=1, value="Severity colors: red = High, yellow = Medium, green = Low. "
            "Counts are live formulas over the Findings sheet.").font = Font(name=FONT, italic=True, size=9)
    if acc_summary is not None:
        r += 1
        ws.cell(row=r, column=1, value="Detection accuracy vs. planted answer key: see the "
                "'Detection Accuracy' sheet.").font = Font(name=FONT, italic=True, size=9)
    for col, w in zip("ABCDEF", (38, 40, 10, 10, 10, 10)):
        ws.column_dimensions[col].width = w

    # --- Findings
    ws = wb.create_sheet("Findings")
    _banner(ws, f"All findings ({n_find})", 0)
    cols = ["finding_id", "issue_type", "severity", "control_id", "application", "account_ref",
            "username", "employee_id", "employee_status", "job_title", "department",
            "store_location", "detail", "last_login_date", "days_since_login", "evidence",
            "recommended_action"]
    _table(ws, findings[cols], 5, sev_col="severity",
           widths={"evidence": 55, "recommended_action": 55, "issue_type": 30, "account_ref": 18})

    # --- Detection accuracy
    if acc_summary is not None:
        ws = wb.create_sheet("Detection Accuracy")
        r = _banner(ws, "Detection accuracy vs. planted_issues.csv", 0)
        last = _table(ws, acc_summary, r)
        ws.cell(row=r, column=7, value="precision").font = Font(name=FONT, bold=True, color="FFFFFF")
        ws.cell(row=r, column=7).fill = HDR_FILL
        ws.cell(row=r, column=8, value="recall").font = Font(name=FONT, bold=True, color="FFFFFF")
        ws.cell(row=r, column=8).fill = HDR_FILL
        for i in range(r + 1, last + 1):
            ws.cell(row=i, column=7, value=f"=IFERROR(D{i}/C{i},0)").number_format = "0.0%"
            ws.cell(row=i, column=8, value=f"=IFERROR(D{i}/B{i},0)").number_format = "0.0%"
        tot = last + 1
        ws.cell(row=tot, column=1, value="Overall").font = Font(name=FONT, bold=True)
        for j in range(2, 7):
            col = get_column_letter(j)
            ws.cell(row=tot, column=j, value=f"=SUM({col}{r + 1}:{col}{last})").font = Font(name=FONT, bold=True)
        ws.cell(row=tot, column=7, value=f"=IFERROR(D{tot}/C{tot},0)").number_format = "0.0%"
        ws.cell(row=tot, column=8, value=f"=IFERROR(D{tot}/B{tot},0)").number_format = "0.0%"
        ws.column_dimensions["G"].width = 12
        ws.column_dimensions["H"].width = 12
        ws.freeze_panes = None
        dr = tot + 3
        ws.cell(row=dr - 1, column=1, value="Item-level comparison").font = Font(name=FONT, bold=True)
        _table(ws, acc_detail[KEY_COLS + ["result"]], dr)
        ws.freeze_panes = None

    # --- Input integrity
    ws = wb.create_sheet("Input Integrity")
    r = _banner(ws, "SHA-256 of input files at the start of this run", 0)
    _table(ws, hashes, r, widths={"sha256": 70})

    # --- Control mapping
    ws = wb.create_sheet("Control Mapping")
    r = _banner(ws, "Finding type to control mapping", 0)
    cm = pd.DataFrame([{"finding_type": t, "control": C.CONTROLS[t][0], "requirement": C.CONTROLS[t][1]}
                       for t in C.FINDING_TYPES])
    _table(ws, cm, r, widths={"requirement": 70})

    # --- Data quality
    ws = wb.create_sheet("Data Quality")
    r = _banner(ws, "HR data-quality warnings (not control violations)", 0)
    wdf = pd.DataFrame(warnings) if warnings else pd.DataFrame([{"warning": "None"}])
    _table(ws, wdf, r)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    digest = sha256_file(path)
    audit(f"analyzer: wrote {path.name}", n_find, digest)
    return digest


def top_risks(findings: pd.DataFrame, n: int = 3) -> pd.DataFrame:
    """Rank finding types by weighted severity (High=3, Medium=2, Low=1)."""
    w = findings["severity"].map({"High": 3, "Medium": 2, "Low": 1})
    score = w.groupby(findings["issue_type"]).sum().sort_values(ascending=False)
    return score.head(n)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--hr", type=Path, default=C.HR_FILE)
    ap.add_argument("--accounts", type=Path, default=C.ACCOUNTS_FILE)
    ap.add_argument("--sod", type=Path, default=C.SOD_FILE)
    args = ap.parse_args()

    audit("analyzer: run started")
    try:
        hr, accts, sod, accounts, findings, warnings, hashes = run_analysis(args.hr, args.accounts, args.sod)
    except ValidationError as exc:
        audit("analyzer: validation FAILED, run stopped")
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    acc_summary, acc_detail = detection_accuracy(findings)
    write_report(findings, accounts, warnings, hashes, acc_summary, acc_detail, C.FINDINGS_XLSX)
    csv_cols = [c for c in findings.columns if c != "control_requirement"]
    write_sanitized_csv(findings[csv_cols], C.FINDINGS_CSV, "analyzer: wrote findings.csv")

    print(f"Accounts reviewed: {len(accounts)}  |  Findings: {len(findings)}")
    print(findings.groupby(["issue_type", "severity"]).size().unstack(fill_value=0).to_string())
    print("\nTop risks (severity-weighted):")
    print(top_risks(findings).to_string())
    if warnings:
        print(f"\nData-quality warnings: {len(warnings)} (see 'Data Quality' sheet)")
    if acc_summary is not None:
        tp, det, pl = (acc_summary[c].sum() for c in ("true_positives", "detected", "planted"))
        print(f"\nDetection accuracy: recall {tp}/{pl} = {tp / pl:.1%}, "
              f"precision {tp}/{det} = {tp / det:.1%}")
        audit("analyzer: detection accuracy computed", int(pl))
    print(f"\nReport: {C.FINDINGS_XLSX.relative_to(C.ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
