"""Process completed certification packets.

Reads packets/completed/*.xlsx, checks each against the issue manifest, and
produces:
  output/campaign_summary.xlsx   completion, revocation, exceptions, reviewers
  output/campaign_decisions.csv  every row's decision (IDs only, no names)
  output/revocations.csv         input for create_tickets.py

Checks performed:
  - Packet integrity: reviewer and row set must match what was issued
  - Incomplete rows: blank or invalid Decision values
  - Missing justification: a flagged row marked Keep with no real reason
  - Rubber-stamping: 100% Keep while flagged items are present
  - Overdue reviewers: not returned, unsigned, or submitted after the deadline

Usage:  python src/process_decisions.py [--review-date YYYY-MM-DD]
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import config as C
from analyzer import _banner, _table
from common import audit, sanitize_value, sha256_file, write_sanitized_csv
from generate_packets import COLUMNS, TABLE_HEADER_ROW

MIN_JUSTIFICATION = 15  # characters; "ok" or "needed" doesn't count
FONT = "Arial"


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def read_packet(path, issued: dict) -> tuple[pd.DataFrame, dict, list[str]]:
    """Return (rows, header, integrity_errors) for one completed packet."""
    errors = []
    wb = load_workbook(path, data_only=True)
    if "Access Review" not in wb.sheetnames or "_meta" not in wb.sheetnames:
        return pd.DataFrame(), {}, ["required sheet missing"]
    meta = {r[0]: r[1] for r in wb["_meta"].iter_rows(values_only=True) if r[0]}
    if meta.get("reviewer_id") != issued["reviewer_id"]:
        errors.append("reviewer_id in packet doesn't match the issued packet")

    ws = wb["Access Review"]
    header = [c.value for c in ws[TABLE_HEADER_ROW]][:len(COLUMNS)]
    if header != COLUMNS:
        return pd.DataFrame(), {}, errors + ["table columns were changed"]
    data = [r[:len(COLUMNS)] for r in ws.iter_rows(min_row=TABLE_HEADER_ROW + 1, values_only=True)
            if any(v not in (None, "") for v in r[:len(COLUMNS)])]
    rows = pd.DataFrame(data, columns=COLUMNS).fillna("")
    rows = rows.astype(str).apply(lambda s: s.str.strip())

    got, expected = set(rows["row_id"]), set(issued["row_ids"])
    if got - expected:
        errors.append(f"{len(got - expected)} row(s) added that were not issued")
    if expected - got:
        errors.append(f"{len(expected - got)} issued row(s) were deleted")
    hdr = {"attestation": str(ws["C11"].value or "").strip(), "submitted": _as_date(ws["C12"].value)}
    return rows, hdr, errors


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--review-date", type=date.fromisoformat,
                    default=C.REVIEW_DEADLINE + timedelta(days=1),
                    help="Date the campaign is evaluated (default: day after the deadline)")
    args = ap.parse_args()
    today, deadline = args.review_date, C.REVIEW_DEADLINE

    if not C.PACKET_MANIFEST.exists():
        print("ERROR: no manifest; run generate_packets.py first", file=sys.stderr)
        return 1
    manifest = json.loads(C.PACKET_MANIFEST.read_text())
    audit("process_decisions: run started", len(manifest["packets"]))

    all_rows, reviewers, exceptions = [], [], []
    for issued in manifest["packets"]:
        rid = issued["reviewer_id"]
        path = C.COMPLETED_DIR / issued["file"]
        rec = {"reviewer_id": rid, "rows": issued["row_count"], "flagged_rows": issued["flagged_rows"],
               "returned": path.exists(), "decided": 0, "kept": 0, "revoked": 0,
               "submitted": "", "status": "", "issues": []}
        if not path.exists():
            rec["status"] = "Overdue - not returned" if today > deadline else "Pending"
            reviewers.append(rec)
            continue
        audit(f"process_decisions: read packet for reviewer {rid}", issued["row_count"], sha256_file(path))
        rows, hdr, integrity = read_packet(path, issued)
        if integrity:
            rec["status"] = "Rejected - integrity check failed"
            rec["issues"] += integrity
            for e in integrity:
                exceptions.append({"reviewer_id": rid, "row_id": "", "employee_id": "",
                                   "exception": "Packet integrity", "detail": e})
            reviewers.append(rec)
            continue

        rows["reviewer_id"] = rid
        valid = rows["Decision"].isin(["Keep", "Revoke"])
        flagged = rows["flags"] != ""
        weak = rows["Justification"].str.len() < MIN_JUSTIFICATION
        for _, r in rows[~valid].iterrows():
            why = "blank" if r["Decision"] == "" else "invalid value"
            exceptions.append({"reviewer_id": rid, "row_id": r["row_id"], "employee_id": r["employee_id"],
                               "exception": "Incomplete row", "detail": f"Decision {why}"})
        for _, r in rows[(rows["Decision"] == "Keep") & flagged & weak].iterrows():
            exceptions.append({"reviewer_id": rid, "row_id": r["row_id"], "employee_id": r["employee_id"],
                               "exception": "Flagged Keep without justification",
                               "detail": f"{r['severity']}: {r['flags']}"})

        rec["decided"] = int(valid.sum())
        rec["kept"] = int((rows["Decision"] == "Keep").sum())
        rec["revoked"] = int((rows["Decision"] == "Revoke").sum())
        rec["submitted"] = hdr["submitted"].isoformat() if hdr["submitted"] else ""

        complete = valid.all() and bool(hdr["attestation"]) and hdr["submitted"] is not None
        rubber = valid.all() and rec["kept"] == len(rows) and flagged.any()
        if rubber:
            rec["issues"].append("Rubber-stamp: 100% Keep with flagged items present")
            exceptions.append({"reviewer_id": rid, "row_id": "", "employee_id": "",
                               "exception": "Rubber-stamping",
                               "detail": f"Kept all {len(rows)} rows incl. {int(flagged.sum())} flagged"})
        if not hdr["attestation"] or hdr["submitted"] is None:
            rec["issues"].append("Attestation or submission date missing")
        if not complete:
            rec["status"] = "Overdue - incomplete" if today > deadline else "In progress"
        elif hdr["submitted"] > deadline:
            rec["status"] = "Completed late"
        else:
            rec["status"] = "Completed on time"
        if rubber:
            rec["status"] += " (rejected: rubber-stamp)"
        all_rows.append(rows)
        reviewers.append(rec)

    rows = pd.concat(all_rows, ignore_index=True) if all_rows else pd.DataFrame(columns=COLUMNS)
    rev = pd.DataFrame(reviewers)
    rev["issues"] = rev["issues"].map("; ".join)
    exc = pd.DataFrame(exceptions, columns=["reviewer_id", "row_id", "employee_id", "exception", "detail"])

    total_rows = int(rev["rows"].sum())
    decided = int(rev["decided"].sum())
    revoked = int(rev["revoked"].sum())
    flagged_total = int(rev["flagged_rows"].sum())
    flagged_revoked = int(((rows["flags"] != "") & (rows["Decision"] == "Revoke")).sum())
    overdue = rev[rev["status"].str.startswith("Overdue") | (rev["status"].str.startswith("Completed late"))]

    # --- Outputs ---------------------------------------------------------
    keep_cols = ["reviewer_id", "row_id", "employee_id", "employee_status", "store_location",
                 "application", "entitlement", "flags", "severity", "controls", "finding_ids",
                 "Decision"]
    decisions = rows[keep_cols].rename(columns={"Decision": "decision"})
    write_sanitized_csv(decisions, C.CAMPAIGN_ROWS_CSV, "process_decisions: wrote campaign_decisions.csv")

    revokes = decisions[decisions["decision"] == "Revoke"].copy()
    # Unflagged revocations still trace to the periodic-review control.
    revokes["controls"] = revokes["controls"].replace("", "NIST 800-53 AC-2 (periodic access review)")
    revokes["severity"] = revokes["severity"].replace("", "Low")
    write_sanitized_csv(revokes, C.REVOCATIONS_CSV, "process_decisions: wrote revocations.csv")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    r = _banner(ws, "Access Certification Campaign Summary", 3)
    metrics = [
        ("Evaluated on", today.isoformat(), None),
        ("Review deadline", deadline.isoformat(), None),
        ("Packets issued", len(rev), None),
        ("Packets returned", int(rev["returned"].sum()), None),
        ("Rows issued", total_rows, None),
        ("Rows decided", decided, None),
        ("Completion rate (rows decided / rows issued)", "=IFERROR(B9/B8,0)", "0.0%"),
        ("Rows revoked", revoked, None),
        ("Revocation rate (revoked / decided)", "=IFERROR(B11/B9,0)", "0.0%"),
        ("Flagged rows issued", flagged_total, None),
        ("Flagged rows revoked", flagged_revoked, None),
        ("Flagged remediation rate", "=IFERROR(B14/B13,0)", "0.0%"),
        ("Overdue or late reviewers", len(overdue), None),
        ("Rubber-stamp reviewers", int(rev["issues"].str.contains("Rubber").sum()), None),
        ("Flagged Keeps missing justification",
         int((exc["exception"] == "Flagged Keep without justification").sum()), None),
        ("Incomplete rows", int((exc["exception"] == "Incomplete row").sum()), None),
    ]
    for label, val, fmt in metrics:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True)
        c = ws.cell(row=r, column=2, value=val)
        c.font = Font(name=FONT)
        if fmt:
            c.number_format = fmt
        r += 1
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 16

    ws = wb.create_sheet("Reviewers")
    r = _banner(ws, "Reviewer status", 0)
    _table(ws, rev, r, widths={"issues": 60, "status": 40})
    bad = PatternFill("solid", start_color="F8CBAD")
    status_col = list(rev.columns).index("status") + 1
    for i in range(r + 1, r + 1 + len(rev)):
        v = str(ws.cell(row=i, column=status_col).value)
        if v.startswith(("Overdue", "Rejected", "Completed late")) or "rubber" in v:
            for j in range(1, len(rev.columns) + 1):
                ws.cell(row=i, column=j).fill = bad

    ws = wb.create_sheet("Exceptions")
    r = _banner(ws, "Exceptions requiring follow-up", 0)
    _table(ws, exc if len(exc) else pd.DataFrame([{"exception": "None"}]), r, widths={"detail": 70})

    ws = wb.create_sheet("Revocations")
    r = _banner(ws, "Approved revocations (sent to remediation)", 0)
    _table(ws, revokes.drop(columns=["decision"]), r)

    C.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(C.CAMPAIGN_XLSX)
    audit("process_decisions: wrote campaign_summary.xlsx", len(rev), sha256_file(C.CAMPAIGN_XLSX))

    # --- Console ---------------------------------------------------------
    print(f"Completion: {decided}/{total_rows} rows ({decided / max(total_rows, 1):.1%})  |  "
          f"Revocation: {revoked}/{decided} ({revoked / max(decided, 1):.1%})  |  "
          f"Flagged remediated: {flagged_revoked}/{flagged_total}")
    print("\nReviewer status:")
    print(rev[["reviewer_id", "rows", "decided", "revoked", "status"]].to_string(index=False))
    print("\nExceptions:")
    print(exc.groupby("exception").size().to_string() if len(exc) else "  none")
    print(f"\nWrote {C.CAMPAIGN_XLSX.relative_to(C.ROOT)} and {C.REVOCATIONS_CSV.relative_to(C.ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
