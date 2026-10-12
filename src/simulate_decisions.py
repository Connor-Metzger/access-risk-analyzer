"""Simulate completed manager packets (for demos and testing ONLY).

Copies each packet to packets/completed/ and fills in decisions the way real
reviewers behave, including the bad behaviors process_decisions.py must
catch. Scenarios are assigned deterministically with the project seed:

  good          revokes risky access, justifies flagged Keeps, signs on time
  rubber_stamp  Keeps 100% of rows even though flagged items are present
  no_justify    Keeps a flagged row without writing a justification
  incomplete    leaves about a third of rows blank and never signs
  late          completes properly but submits after the deadline
  not_returned  never sends the packet back

Usage:  python src/simulate_decisions.py
"""
from __future__ import annotations

import json
import random
import shutil
import sys
from datetime import timedelta

from openpyxl import load_workbook

import config as C
from common import audit, sha256_file
from generate_packets import COLUMNS, ORPHAN_REVIEWER, TABLE_HEADER_ROW

# Seeded PRNG is intentional: reproducible demo decisions, not security use.
rng = random.Random(C.SEED + 1)  # nosec B311
DEC = COLUMNS.index("Decision") + 1
JUS = COLUMNS.index("Justification") + 1
FLAGS = COLUMNS.index("flags") + 1
SEV = COLUMNS.index("severity") + 1

KEEP_REASONS = [
    "Seasonal employee returning for holiday season; confirmed with HR",
    "Acting lead on night shift through 2026-11-30; time-bound, HR ticket on file",
    "Business need confirmed; secondary review by district manager in place",
]


def assign_scenarios(packets: list[dict]) -> dict[str, str]:
    flagged = sorted(p["reviewer_id"] for p in packets
                     if p["flagged_rows"] and p["reviewer_id"] != ORPHAN_REVIEWER)
    clean = sorted(p["reviewer_id"] for p in packets
                   if not p["flagged_rows"] and p["reviewer_id"] != ORPHAN_REVIEWER)
    rng.shuffle(flagged)
    rng.shuffle(clean)
    scen = {r: "good" for r in flagged + clean}
    scen[ORPHAN_REVIEWER] = "good"
    for r, s in zip(flagged, ["rubber_stamp", "no_justify", "incomplete", "late"]):
        scen[r] = s
    if clean:
        scen[clean[0]] = "not_returned"
    return scen


def decide(scenario: str, flags: str, severity: str, first_flag_kept: list) -> tuple[str, str]:
    if scenario == "rubber_stamp":
        return "Keep", ""
    if not flags:
        return ("Revoke", "No longer needed for current role") if rng.random() < 0.03 else ("Keep", "")
    if severity == "High":
        return "Revoke", ""
    # Medium/Low flagged rows: some are legitimately kept with a reason.
    if scenario == "no_justify" and not first_flag_kept:
        first_flag_kept.append(True)
        return "Keep", ""
    if rng.random() < 0.4:
        return "Keep", rng.choice(KEEP_REASONS)
    return "Revoke", ""


def main() -> int:
    if not C.PACKET_MANIFEST.exists():
        print("ERROR: run generate_packets.py first", file=sys.stderr)
        return 1
    manifest = json.loads(C.PACKET_MANIFEST.read_text())
    scenarios = assign_scenarios(manifest["packets"])
    if C.COMPLETED_DIR.exists():
        shutil.rmtree(C.COMPLETED_DIR)
    C.COMPLETED_DIR.mkdir(parents=True)

    for p in manifest["packets"]:
        rid, scen = p["reviewer_id"], scenarios[p["reviewer_id"]]
        if scen == "not_returned":
            print(f"  {rid:<12} {scen}")
            continue
        wb = load_workbook(C.PACKETS_DIR / p["file"])
        ws = wb["Access Review"]
        kept = []
        rows = list(range(TABLE_HEADER_ROW + 1, TABLE_HEADER_ROW + 1 + p["row_count"]))
        blank = set(rng.sample(rows, k=len(rows) // 3)) if scen == "incomplete" else set()
        for r in rows:
            if r in blank:
                continue
            d, j = decide(scen, ws.cell(r, FLAGS).value or "", ws.cell(r, SEV).value or "", kept)
            ws.cell(r, DEC).value = d
            ws.cell(r, JUS).value = j or None
        if scen != "incomplete":
            ws["C11"] = "Signed electronically" if rid == ORPHAN_REVIEWER else f"Reviewer {rid}"
            offset = rng.randint(3, 6) if scen == "late" else -rng.randint(0, 10)
            ws["C12"] = (C.REVIEW_DEADLINE + timedelta(days=offset)).isoformat()
        out = C.COMPLETED_DIR / p["file"]
        wb.save(out)
        audit(f"simulate_decisions: completed packet for reviewer {rid} ({scen})",
              p["row_count"], sha256_file(out))
        print(f"  {rid:<12} {scen}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
